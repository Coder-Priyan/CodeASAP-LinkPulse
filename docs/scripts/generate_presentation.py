"""
LinkPulse - Round 1 Presentation Generator (PowerPoint .pptx)
Innovista 2026 | Graphic Era Hill University (GEHU), Haldwani Campus
Generates a 10-slide 16:9 presentation with editable shapes, text boxes, and original mascots.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Color Palette Constants
COLOR_BG = RGBColor(255, 255, 255)            # Clean White
COLOR_DARK = RGBColor(15, 23, 42)             # Deep Navy / Charcoal (#0F172A)
COLOR_TEXT_PRIMARY = RGBColor(15, 23, 42)     # Deep Navy
COLOR_TEXT_MUTED = RGBColor(71, 85, 105)      # Slate Gray (#475569)
COLOR_CYAN = RGBColor(2, 132, 199)            # Vibrant Cyan/Blue (#0284C7)
COLOR_CYAN_LIGHT = RGBColor(224, 242, 254)    # Light Cyan (#E0F2FE)
COLOR_YELLOW = RGBColor(254, 240, 138)        # Warm Card Yellow (#FEF08A)
COLOR_YELLOW_BORDER = RGBColor(234, 179, 8)   # Gold Border (#EAB308)
COLOR_CARD_BG = RGBColor(248, 250, 252)       # Off White (#F8FAFC)
COLOR_CARD_BORDER = RGBColor(15, 23, 42)      # Dark Border (#0F172A)
COLOR_GREEN_BG = RGBColor(220, 252, 231)      # Soft Green (#DCFCE7)
COLOR_GREEN_TEXT = RGBColor(22, 101, 52)      # Forest Green (#166534)
COLOR_RED_TEXT = RGBColor(225, 29, 72)        # Rose Red (#E11D48)

IMG_DIR = r"C:\Users\hp\.gemini\antigravity-ide\brain\aeb5f0e2-5c19-449f-b7fe-6bd5d39f5b7f\scratch\extracted_images"

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Helper: Add Slide Header Badge
    def add_header_badge(slide, text, width=3.4, height=0.6, left=0.8, top=0.5):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CYAN
        shape.line.color.rgb = COLOR_DARK
        shape.line.width = Pt(2)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        return shape

    # Helper: Add Rounded Card
    def add_card(slide, left, top, width, height, fill_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER, border_width=2):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        return card

    # Helper: Add Callout Banner at bottom
    def add_bottom_callout(slide, title, text, left=0.8, top=6.25, width=11.733, height=0.8):
        banner = add_card(slide, left, top, width, height, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
        tf = banner.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.text = title + ": "
        run = p.add_run()
        run.text = text
        run.font.bold = False
        return banner

    # =========================================================================
    # SLIDE 1: Project Introduction
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    
    # Event Badge
    add_header_badge(s1, "INNOVISTA 2026 | ROUND 1: IDEATION", width=4.0, height=0.55, left=0.8, top=0.5)

    # Main Title
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(7.5), Inches(1.2))
    p1 = tb_title.text_frame.paragraphs[0]
    p1.text = "LINKPULSE"
    p1.font.name = "Arial Black"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_CYAN

    # Subtitle
    p2 = tb_title.text_frame.add_paragraph()
    p2.text = "URL Shortener & Telemetry Platform"
    p2.font.name = "Arial"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(4)

    # Team Card
    team_card = add_card(s1, 0.8, 3.0, 7.8, 3.8, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tf_team = team_card.text_frame
    tf_team.margin_left = Inches(0.35)
    tf_team.margin_top = Inches(0.25)
    tf_team.margin_right = Inches(0.35)

    p_t1 = tf_team.paragraphs[0]
    p_t1.text = "TEAM NAME"
    p_t1.font.name = "Arial Black"
    p_t1.font.size = Pt(12)
    p_t1.font.color.rgb = COLOR_DARK

    p_t2 = tf_team.add_paragraph()
    p_t2.text = "[TEAM NAME]"
    p_t2.font.name = "Arial Black"
    p_t2.font.size = Pt(22)
    p_t2.font.color.rgb = COLOR_DARK
    p_t2.space_after = Pt(10)

    p_t3 = tf_team.add_paragraph()
    p_t3.text = "TEAM MEMBERS & ROLES"
    p_t3.font.name = "Arial Black"
    p_t3.font.size = Pt(11)
    p_t3.font.color.rgb = COLOR_DARK
    p_t3.space_after = Pt(4)

    members = [
        ("Abhishek Dhapola", "Frontend and User Interface"),
        ("Sandeep Singh Mehta", "Database Architecture and Data Modeling"),
        ("Aditya Mehra", "Testing, Quality Assurance and Security"),
        ("Deepak Kumar Kabi", "Technical Documentation and Presentation"),
        ("Priyanshu Dangi", "Team Lead, Backend and System Design")
    ]
    for name, role in members:
        p_m = tf_team.add_paragraph()
        p_m.text = f"•  {name} — {role}"
        p_m.font.name = "Arial"
        p_m.font.size = Pt(11)
        p_m.font.color.rgb = COLOR_DARK
        p_m.space_before = Pt(2)

    # Speech Bubble "Hello!"
    bubble = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(0.8), Inches(1.6), Inches(0.6))
    bubble.fill.solid()
    bubble.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bubble.line.color.rgb = COLOR_DARK
    bubble.line.width = Pt(2)
    bubble.text_frame.paragraphs[0].text = "HELLO!"
    bubble.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    bubble.text_frame.paragraphs[0].font.name = "Arial Black"
    bubble.text_frame.paragraphs[0].font.size = Pt(13)
    bubble.text_frame.paragraphs[0].font.color.rgb = COLOR_DARK

    # Mascot Image
    img_mascot_s1 = os.path.join(IMG_DIR, "Im25.png")
    if os.path.exists(img_mascot_s1):
        s1.shapes.add_picture(img_mascot_s1, Inches(9.2), Inches(1.3), width=Inches(3.4))

    # =========================================================================
    # SLIDE 2: Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header_badge(s2, "PROBLEM STATEMENT", width=3.4, height=0.55, left=0.8, top=0.5)

    # Subtitle
    tb_s2 = s2.shapes.add_textbox(Inches(4.4), Inches(0.5), Inches(8.0), Inches(0.6))
    p_s2_sub = tb_s2.text_frame.paragraphs[0]
    p_s2_sub.text = "Why Standard URL Shorteners Break Under Concurrent Traffic"
    p_s2_sub.font.name = "Arial"
    p_s2_sub.font.size = Pt(15)
    p_s2_sub.font.bold = True
    p_s2_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Box 1: 1 Click Left
    b1 = add_card(s2, 0.8, 1.35, 2.3, 3.2, fill_color=RGBColor(255, 255, 255))
    tf1 = b1.text_frame
    tf1.margin_top = Inches(0.2)
    p = tf1.paragraphs[0]
    p.text = "1. ONE CLICK LEFT"
    p.font.name = "Arial Black"
    p.font.size = Pt(12)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = COLOR_DARK

    p_badge1 = tf1.add_paragraph()
    p_badge1.text = "short link"
    p_badge1.font.name = "Arial"
    p_badge1.font.size = Pt(10)
    p_badge1.alignment = PP_ALIGN.CENTER
    p_badge1.space_before = Pt(8)

    p_num1 = tf1.add_paragraph()
    p_num1.text = "1"
    p_num1.font.name = "Arial Black"
    p_num1.font.size = Pt(56)
    p_num1.font.color.rgb = COLOR_RED_TEXT
    p_num1.alignment = PP_ALIGN.CENTER

    p_lbl1 = tf1.add_paragraph()
    p_lbl1.text = "click remaining before closure"
    p_lbl1.font.name = "Arial"
    p_lbl1.font.size = Pt(10)
    p_lbl1.font.color.rgb = COLOR_TEXT_MUTED
    p_lbl1.alignment = PP_ALIGN.CENTER

    # Box 2: Same Moment
    b2 = add_card(s2, 3.3, 1.35, 2.5, 3.2, fill_color=RGBColor(255, 255, 255))
    tf2 = b2.text_frame
    tf2.margin_top = Inches(0.2)
    p = tf2.paragraphs[0]
    p.text = "2. SIMULTANEOUS CLICKS"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = COLOR_DARK

    # Characters in Box 2
    img_c1 = os.path.join(IMG_DIR, "Im56.png")
    img_c2 = os.path.join(IMG_DIR, "Im58.png")
    if os.path.exists(img_c1) and os.path.exists(img_c2):
        s2.shapes.add_picture(img_c1, Inches(3.5), Inches(2.1), width=Inches(0.95))
        s2.shapes.add_picture(img_c2, Inches(4.55), Inches(2.1), width=Inches(0.95))

    p_lbl2 = tf2.add_paragraph()
    p_lbl2.text = "two visitors click at the exact same moment"
    p_lbl2.font.name = "Arial"
    p_lbl2.font.size = Pt(10)
    p_lbl2.font.color.rgb = COLOR_TEXT_MUTED
    p_lbl2.alignment = PP_ALIGN.CENTER
    p_lbl2.space_before = Pt(110)

    # Box 3: Both Get In
    b3 = add_card(s2, 6.0, 1.35, 2.4, 3.2, fill_color=RGBColor(255, 255, 255))
    tf3 = b3.text_frame
    tf3.margin_top = Inches(0.2)
    p = tf3.paragraphs[0]
    p.text = "3. LIMIT BROKEN"
    p.font.name = "Arial Black"
    p.font.size = Pt(12)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = COLOR_DARK

    p_num3 = tf3.add_paragraph()
    p_num3.text = "2"
    p_num3.font.name = "Arial Black"
    p_num3.font.size = Pt(56)
    p_num3.font.color.rgb = COLOR_RED_TEXT
    p_num3.alignment = PP_ALIGN.CENTER
    p_num3.space_before = Pt(10)

    p_lbl3 = tf3.add_paragraph()
    p_lbl3.text = "visitors allowed through!\nLimit over-allocated"
    p_lbl3.font.name = "Arial"
    p_lbl3.font.size = Pt(10)
    p_lbl3.font.color.rgb = COLOR_TEXT_MUTED
    p_lbl3.alignment = PP_ALIGN.CENTER

    # Box 4: The Problem in Short (Right side)
    b4 = add_card(s2, 8.6, 1.35, 3.933, 4.7, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tf4 = b4.text_frame
    tf4.margin_left = Inches(0.25)
    tf4.margin_top = Inches(0.25)
    tf4.margin_right = Inches(0.25)
    p = tf4.paragraphs[0]
    p.text = "THE PROBLEM IN SHORT"
    p.font.name = "Arial Black"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK

    short_points = [
        "Basic shorteners only redirect URLs; they lack rule engines, campaign grouping, and quick takedown tools.",
        "Limited-capacity promotions (VIP passes, flash sales) easily exceed limits when traffic spikes.",
        "Synchronous tracking delays redirects and creates a sluggish user experience.",
        "Raw click counters fail to distinguish genuine visitors from search bots, web scrapers, and crawlers."
    ]
    for pt in short_points:
        p_pt = tf4.add_paragraph()
        p_pt.text = f"•  {pt}"
        p_pt.font.name = "Arial"
        p_pt.font.size = Pt(11)
        p_pt.font.color.rgb = COLOR_DARK
        p_pt.space_before = Pt(8)

    # Box 5: Why It Happens (Bottom left)
    b5 = add_card(s2, 0.8, 4.75, 7.6, 1.3, fill_color=COLOR_CARD_BG, border_color=COLOR_DARK, border_width=2)
    tf5 = b5.text_frame
    tf5.margin_left = Inches(0.25)
    tf5.margin_top = Inches(0.15)
    p = tf5.paragraphs[0]
    p.text = "WHY IT HAPPENS & HOW LINKPULSE SOLVES IT"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_CYAN

    p_sol = tf5.add_paragraph()
    p_sol.text = "Standard applications read the click count into server memory, check the limit, and update the database later. Under concurrent visits, both requests read the old count and both pass. LinkPulse checks and increments the counter in a single atomic SQL step, stopping the extra visitor immediately."
    p_sol.font.name = "Arial"
    p_sol.font.size = Pt(10)
    p_sol.font.color.rgb = COLOR_DARK
    p_sol.space_before = Pt(3)

    # =========================================================================
    # SLIDE 3: Solution Overview
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header_badge(s3, "SOLUTION OVERVIEW", width=3.4, height=0.55, left=0.8, top=0.5)

    # Diagram Entities
    # User Box
    add_card(s3, 0.8, 1.4, 2.5, 0.9, fill_color=RGBColor(255, 255, 255))
    tb_u = s3.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(2.5), Inches(0.9))
    tb_u.text_frame.paragraphs[0].text = "USER (OWNER)"
    tb_u.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_u.text_frame.paragraphs[0].font.size = Pt(11)
    tb_u.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    p_u_sub = tb_u.text_frame.add_paragraph()
    p_u_sub.text = "Creates & manages links"
    p_u_sub.font.name = "Arial"
    p_u_sub.font.size = Pt(9.5)
    p_u_sub.alignment = PP_ALIGN.CENTER

    # Campaigns Box
    add_card(s3, 4.4, 1.4, 2.8, 0.9, fill_color=RGBColor(255, 255, 255))
    tb_c = s3.shapes.add_textbox(Inches(4.4), Inches(1.4), Inches(2.8), Inches(0.9))
    tb_c.text_frame.paragraphs[0].text = "CAMPAIGNS & TAGS"
    tb_c.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_c.text_frame.paragraphs[0].font.size = Pt(11)
    tb_c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    p_c_sub = tb_c.text_frame.add_paragraph()
    p_c_sub.text = "Organizes & groups links"
    p_c_sub.font.name = "Arial"
    p_c_sub.font.size = Pt(9.5)
    p_c_sub.alignment = PP_ALIGN.CENTER

    # LINK Central Box
    add_card(s3, 2.2, 2.8, 3.2, 2.8, fill_color=COLOR_CYAN_LIGHT, border_color=COLOR_DARK, border_width=2.5)
    tb_l = s3.shapes.add_textbox(Inches(2.2), Inches(2.85), Inches(3.2), Inches(2.7))
    p = tb_l.text_frame.paragraphs[0]
    p.text = "CENTRAL LINK RECORD"
    p.font.name = "Arial Black"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    link_fields = [
        "Short Code (Base62 / Alias)",
        "Destination URL (Validated)",
        "Active Status & Deactivation Reason",
        "Expiry Date & Click Limit",
        "Passcode Hash (bcrypt)",
        "Atomic Click Counter"
    ]
    for lf in link_fields:
        p_lf = tb_l.text_frame.add_paragraph()
        p_lf.text = f"•  {lf}"
        p_lf.font.name = "Arial"
        p_lf.font.size = Pt(9.5)
        p_lf.font.color.rgb = COLOR_DARK
        p_lf.space_before = Pt(2)

    # ClickEvent Box
    add_card(s3, 5.9, 3.4, 2.3, 1.4, fill_color=RGBColor(255, 255, 255))
    tb_ce = s3.shapes.add_textbox(Inches(5.9), Inches(3.45), Inches(2.3), Inches(1.3))
    tb_ce.text_frame.paragraphs[0].text = "CLICKEVENT"
    tb_ce.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_ce.text_frame.paragraphs[0].font.size = Pt(11)
    tb_ce.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    ce_pts = ["Recorded per human visit", "Device, browser, OS, referrer", "Salted visitor hash"]
    for cp in ce_pts:
        p_cp = tb_ce.text_frame.add_paragraph()
        p_cp.text = f"• {cp}"
        p_cp.font.name = "Arial"
        p_cp.font.size = Pt(8.5)
        p_cp.font.color.rgb = COLOR_DARK

    # Dashboard Box
    add_card(s3, 2.4, 5.9, 3.8, 0.85, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tb_db = s3.shapes.add_textbox(Inches(2.4), Inches(5.9), Inches(3.8), Inches(0.85))
    p = tb_db.text_frame.paragraphs[0]
    p.text = "LIVE ANALYTICS DASHBOARD"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.alignment = PP_ALIGN.CENTER
    p_db = tb_db.text_frame.add_paragraph()
    p_db.text = "Aggregates clicks, unique visitors, trends & alerts"
    p_db.font.name = "Arial"
    p_db.font.size = Pt(9)
    p_db.alignment = PP_ALIGN.CENTER

    # Mascot on Left
    img_mascot_s3 = os.path.join(IMG_DIR, "Im134.png")
    if os.path.exists(img_mascot_s3):
        s3.shapes.add_picture(img_mascot_s3, Inches(0.6), Inches(4.3), width=Inches(1.5))

    # Right Box: Core Engineering Decisions
    b_dec = add_card(s3, 8.5, 1.4, 4.033, 5.35, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    tf_dec = b_dec.text_frame
    tf_dec.margin_left = Inches(0.3)
    tf_dec.margin_top = Inches(0.3)
    tf_dec.margin_right = Inches(0.3)

    p = tf_dec.paragraphs[0]
    p.text = "CORE ENGINEERING DECISIONS"
    p.font.name = "Arial Black"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK

    decisions = [
        ("Central Link Hub", "The Link entity is the central record connecting owners, campaigns, redirect policies, and telemetry logs."),
        ("Strict Rule Hierarchy", "Every visit evaluates rules in a fixed order (Disabled → Expired → Click Limit → Password) for 100% predictable outcomes."),
        ("Non-Blocking Telemetry", "Redirects respond immediately with HTTP 307; click logging is dispatched asynchronously in the background via Next.js after()."),
        ("Historical Integrity", "Deactivating a link or deleting a campaign decouples associations but preserves all past click history for audit and reporting.")
    ]
    for d_title, d_desc in decisions:
        p_dt = tf_dec.add_paragraph()
        p_dt.text = f"•  {d_title}: "
        p_dt.font.name = "Arial"
        p_dt.font.size = Pt(11)
        p_dt.font.bold = True
        p_dt.font.color.rgb = COLOR_DARK
        p_dt.space_before = Pt(10)
        run = p_dt.add_run()
        run.text = d_desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Tech Stack
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header_badge(s4, "TECH STACK", width=3.0, height=0.55, left=0.8, top=0.5)

    # Table of Tech Stack
    rows, cols = 6, 3
    table_shape = s4.shapes.add_table(rows, cols, Inches(0.8), Inches(1.3), Inches(11.733), Inches(4.7))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(4.3)
    tbl.columns[2].width = Inches(5.233)

    headers = ["LAYER", "WHAT WE USE", "WHY WE CHOSE IT"]
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CYAN
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = "Arial Black"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(255, 255, 255)

    stack_data = [
        ("App & API", "Next.js 15 (App Router) + TypeScript", "Combines UI pages, server components, and API route handlers in a unified type-safe codebase."),
        ("Data Access", "Prisma ORM 6", "Declarative schema provides a single source of truth with fully typed queries and migrations."),
        ("Database", "PostgreSQL on Supabase", "Production relational database supporting atomic conditional queries, indexes, and complex joins."),
        ("Sessions", "Stateless JWT in httpOnly Cookies", "Secure, signed authentication cookies that prevent client-side JavaScript theft (anti-XSS)."),
        ("Cloud Hosting", "Vercel (Planned Target)", "Serverless deployment platform optimized natively for Next.js 15 App Router and edge execution.")
    ]

    for r, (layer, tech, why) in enumerate(stack_data, start=1):
        # Layer cell
        c0 = tbl.cell(r, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = RGBColor(241, 245, 249) if r % 2 == 0 else RGBColor(255, 255, 255)
        p0 = c0.text_frame.paragraphs[0]
        p0.text = layer
        p0.font.name = "Arial"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_DARK

        # Tech cell
        c1 = tbl.cell(r, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = RGBColor(241, 245, 249) if r % 2 == 0 else RGBColor(255, 255, 255)
        p1 = c1.text_frame.paragraphs[0]
        p1.text = tech
        p1.font.name = "Arial"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_CYAN

        # Why cell
        c2 = tbl.cell(r, 2)
        c2.fill.solid()
        c2.fill.fore_color.rgb = RGBColor(241, 245, 249) if r % 2 == 0 else RGBColor(255, 255, 255)
        p2 = c2.text_frame.paragraphs[0]
        p2.text = why
        p2.font.name = "Arial"
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Callout
    add_bottom_callout(s4, "Key Technical Decisions", 
        "HTTP 307 temporary redirects with no-store cache control (prevents browser caching so every visit is tracked), atomic conditional SQL (prevents limit overshoots), and salted SHA-256 visitor hashing (no raw IPs stored).")

    # =========================================================================
    # SLIDE 5: Implementation Status
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header_badge(s5, "IMPLEMENTATION STATUS", width=4.0, height=0.55, left=0.8, top=0.5)

    # Architecture Pipeline on top
    pipe_box = add_card(s5, 0.8, 1.25, 10.2, 0.75, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    tf_pipe = pipe_box.text_frame
    tf_pipe.margin_top = Inches(0.12)
    p = tf_pipe.paragraphs[0]
    p.text = "Route Handlers (Thin APIs)  ➔  Services Layer (Rules & Telemetry)  ➔  Prisma ORM  ➔  PostgreSQL"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Mascot with wrench
    img_mascot_s5 = os.path.join(IMG_DIR, "Im421.png")
    if os.path.exists(img_mascot_s5):
        s5.shapes.add_picture(img_mascot_s5, Inches(11.3), Inches(0.5), width=Inches(1.2))

    # 3 Columns
    col_w = 3.65
    gap = 0.39

    # Col 1: Built and Working
    col1 = add_card(s5, 0.8, 2.2, col_w, 4.8, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    b1_badge = add_card(s5, 0.95, 2.35, col_w - 0.3, 0.45, fill_color=COLOR_GREEN_BG, border_color=COLOR_GREEN_TEXT, border_width=1.5)
    b1_badge.text_frame.paragraphs[0].text = "BUILT & VERIFIED (MILESTONES 1–3)"
    b1_badge.text_frame.paragraphs[0].font.name = "Arial Black"
    b1_badge.text_frame.paragraphs[0].font.size = Pt(10)
    b1_badge.text_frame.paragraphs[0].font.color.rgb = COLOR_GREEN_TEXT
    b1_badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    tf_c1 = s5.shapes.add_textbox(Inches(0.95), Inches(2.9), Inches(col_w - 0.3), Inches(4.0)).text_frame
    built_items = [
        "Role-based auth (USER & ADMIN)",
        "2-step link creation & alias validation",
        "Rule-driven redirect pipeline",
        "Expiry, atomic click limits, password hash",
        "Bot filtering & unique visitor hashing",
        "Campaigns & tags (with safe cascade)",
        "Live analytics dashboard (7d, 30d, all-time)",
        "Device, browser, referrer breakdowns",
        "Security: password rate limits & salt isolation",
        "Automated test suite (27 Vitest tests passing)"
    ]
    for idx, item in enumerate(built_items):
        p = tf_c1.paragraphs[0] if idx == 0 else tf_c1.add_paragraph()
        p.text = f"✓  {item}"
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(3)

    # Col 2: In Progress & Next for Round 2
    col2 = add_card(s5, 0.8 + col_w + gap, 2.2, col_w, 4.8, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    b2_badge = add_card(s5, 0.95 + col_w + gap, 2.35, col_w - 0.3, 0.45, fill_color=COLOR_CYAN_LIGHT, border_color=COLOR_CYAN, border_width=1.5)
    b2_badge.text_frame.paragraphs[0].text = "IN PROGRESS & NEXT FOR ROUND 2"
    b2_badge.text_frame.paragraphs[0].font.name = "Arial Black"
    b2_badge.text_frame.paragraphs[0].font.size = Pt(10)
    b2_badge.text_frame.paragraphs[0].font.color.rgb = COLOR_CYAN
    b2_badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    tf_c2 = s5.shapes.add_textbox(Inches(0.95 + col_w + gap), Inches(2.9), Inches(col_w - 0.3), Inches(4.0)).text_frame
    next_items = [
        "Cloud production deployment on Vercel",
        "QR code generation with dynamic scan tracking",
        "Bulk CSV link import with duplicate validation",
        "Admin domain blocklist for malware & phishing",
        "Refining E2E browser test flows",
        "Enhanced mobile responsiveness audit"
    ]
    for idx, item in enumerate(next_items):
        p = tf_c2.paragraphs[0] if idx == 0 else tf_c2.add_paragraph()
        p.text = f"➔  {item}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(8)

    # Col 3: Future Roadmap
    col3 = add_card(s5, 0.8 + (col_w + gap) * 2, 2.2, col_w, 4.8, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    b3_badge = add_card(s5, 0.95 + (col_w + gap) * 2, 2.35, col_w - 0.3, 0.45, fill_color=RGBColor(241, 245, 249), border_color=COLOR_TEXT_MUTED, border_width=1.5)
    b3_badge.text_frame.paragraphs[0].text = "FUTURE ROADMAP"
    b3_badge.text_frame.paragraphs[0].font.name = "Arial Black"
    b3_badge.text_frame.paragraphs[0].font.size = Pt(10)
    b3_badge.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_MUTED
    b3_badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    tf_c3 = s5.shapes.add_textbox(Inches(0.95 + (col_w + gap) * 2), Inches(2.9), Inches(col_w - 0.3), Inches(4.0)).text_frame
    future_items = [
        "Custom branded domains (CNAME mapping)",
        "Device-specific redirect targeting (iOS/Android)",
        "Public developer REST API with API keys",
        "Link-in-bio customizable landing pages",
        "Real-time webhook notifications for alerts",
        "Edge GeoIP location telemetry"
    ]
    for idx, item in enumerate(future_items):
        p = tf_c3.paragraphs[0] if idx == 0 else tf_c3.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 6: Application Architecture
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header_badge(s6, "APPLICATION ARCHITECTURE", width=4.2, height=0.55, left=0.8, top=0.5)

    # Left: User Roles Column
    add_card(s6, 0.8, 1.3, 2.2, 1.25, fill_color=RGBColor(255, 255, 255))
    tb_vis = s6.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(2.2), Inches(1.15))
    tb_vis.text_frame.paragraphs[0].text = "VISITOR"
    tb_vis.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_vis.text_frame.paragraphs[0].font.size = Pt(11)
    tb_vis.text_frame.paragraphs[0].font.color.rgb = COLOR_DARK
    p = tb_vis.text_frame.add_paragraph()
    p.text = "Opens short links & enters passcodes"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s6, 0.8, 2.7, 2.2, 1.25, fill_color=RGBColor(255, 255, 255))
    tb_usr = s6.shapes.add_textbox(Inches(0.8), Inches(2.75), Inches(2.2), Inches(1.15))
    tb_usr.text_frame.paragraphs[0].text = "USER"
    tb_usr.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_usr.text_frame.paragraphs[0].font.size = Pt(11)
    tb_usr.text_frame.paragraphs[0].font.color.rgb = COLOR_CYAN
    p = tb_usr.text_frame.add_paragraph()
    p.text = "Manages links, campaigns & own analytics"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s6, 0.8, 4.1, 2.2, 1.25, fill_color=RGBColor(255, 255, 255))
    tb_adm = s6.shapes.add_textbox(Inches(0.8), Inches(4.15), Inches(2.2), Inches(1.15))
    tb_adm.text_frame.paragraphs[0].text = "ADMINISTRATOR"
    tb_adm.text_frame.paragraphs[0].font.name = "Arial Black"
    tb_adm.text_frame.paragraphs[0].font.size = Pt(11)
    tb_adm.text_frame.paragraphs[0].font.color.rgb = COLOR_RED_TEXT
    p = tb_adm.text_frame.add_paragraph()
    p.text = "Platform oversight & link deactivation"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    tb_sess = s6.shapes.add_textbox(Inches(0.8), Inches(5.45), Inches(2.2), Inches(0.6))
    p_sess = tb_sess.text_frame.paragraphs[0]
    p_sess.text = "Session: Signed JWT in httpOnly Cookie"
    p_sess.font.name = "Arial"
    p_sess.font.size = Pt(9)
    p_sess.font.color.rgb = COLOR_TEXT_MUTED
    p_sess.font.bold = True

    # Middle: Next.js 15 App Core
    add_card(s6, 3.3, 1.3, 5.4, 4.7, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2.5)
    tb_core = s6.shapes.add_textbox(Inches(3.3), Inches(1.4), Inches(5.4), Inches(0.6))
    p = tb_core.text_frame.paragraphs[0]
    p.text = "NEXT.JS 15 (APP ROUTER) + TYPESCRIPT"
    p.font.name = "Arial Black"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    # Presentation Box
    add_card(s6, 3.5, 2.1, 5.0, 0.75, fill_color=RGBColor(241, 245, 249))
    tb_pres = s6.shapes.add_textbox(Inches(3.5), Inches(2.15), Inches(5.0), Inches(0.65))
    p = tb_pres.text_frame.paragraphs[0]
    p.text = "Presentation Layer (Pages & UI Components)"
    p.font.name = "Arial Black"
    p.font.size = Pt(10)
    p.alignment = PP_ALIGN.CENTER
    p_sub = tb_pres.text_frame.add_paragraph()
    p_sub.text = "Landing Hero, 2-Step Creator, Analytics Dashboard, Admin Console"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(8.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    # Route Handlers Box
    add_card(s6, 3.5, 3.0, 5.0, 0.65, fill_color=RGBColor(241, 245, 249))
    tb_rh = s6.shapes.add_textbox(Inches(3.5), Inches(3.05), Inches(5.0), Inches(0.55))
    p = tb_rh.text_frame.paragraphs[0]
    p.text = "Thin API Route Handlers (Validation & Guards)"
    p.font.name = "Arial Black"
    p.font.size = Pt(10)
    p.alignment = PP_ALIGN.CENTER
    p_sub = tb_rh.text_frame.add_paragraph()
    p_sub.text = "/[code] Redirect, /api/links, /api/campaigns, /api/analytics"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(8.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    # Services Layer
    add_card(s6, 3.5, 3.8, 5.0, 1.95, fill_color=COLOR_CYAN_LIGHT, border_color=COLOR_CYAN, border_width=1.5)
    tb_srv = s6.shapes.add_textbox(Inches(3.5), Inches(3.85), Inches(5.0), Inches(1.85))
    p = tb_srv.text_frame.paragraphs[0]
    p.text = "SERVICES LAYER (DECOUPLED BUSINESS LOGIC)"
    p.font.name = "Arial Black"
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    services_list = [
        ("Auth & Roles Guard", "User registration, password verification & session guards"),
        ("Link & Rules Service", "Base62 code generation, rule evaluation, atomic limits"),
        ("Click Telemetry Service", "Bot detection, metadata extraction & privacy hash"),
        ("Analytics Engine", "Database-level aggregations across periods & campaigns")
    ]
    for s_name, s_desc in services_list:
        p_s = tb_srv.text_frame.add_paragraph()
        p_s.text = f"•  {s_name}: "
        p_s.font.name = "Arial"
        p_s.font.size = Pt(8.5)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_DARK
        p_s.space_before = Pt(2)
        run = p_s.add_run()
        run.text = s_desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Database Layer
    add_card(s6, 9.0, 1.3, 3.533, 4.7, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2.5)
    tb_db_title = s6.shapes.add_textbox(Inches(9.0), Inches(1.4), Inches(3.533), Inches(0.8))
    p = tb_db_title.text_frame.paragraphs[0]
    p.text = "DATA LAYER"
    p.font.name = "Arial Black"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p.alignment = PP_ALIGN.CENTER
    p_db_sub = tb_db_title.text_frame.add_paragraph()
    p_db_sub.text = "Prisma ORM ➔ PostgreSQL (Supabase)"
    p_db_sub.font.name = "Arial"
    p_db_sub.font.size = Pt(10)
    p_db_sub.font.bold = True
    p_db_sub.font.color.rgb = COLOR_CYAN
    p_db_sub.alignment = PP_ALIGN.CENTER

    schema_tables = [
        ("User", "Accounts, roles, hashed passwords"),
        ("Link", "Codes, destinations, rules, atomic counters"),
        ("ClickEvent", "Timestamp, device, browser, OS, visitorHash"),
        ("Campaign", "Marketing groupings with safe delete"),
        ("Tag & LinkTag", "Many-to-many labels for filtering"),
        ("PasswordAttempt", "Salted IP hashes for sliding lockout")
    ]
    tb_tbls = s6.shapes.add_textbox(Inches(9.2), Inches(2.3), Inches(3.133), Inches(3.5))
    for idx, (t_name, t_desc) in enumerate(schema_tables):
        p = tb_tbls.text_frame.paragraphs[0] if idx == 0 else tb_tbls.text_frame.add_paragraph()
        p.text = f"•  {t_name}: "
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(4)
        run = p.add_run()
        run.text = t_desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Callout
    add_bottom_callout(s6, "High-Performance Redirect Flow", 
        "Visitor ➔ Thin Route Handler ➔ Business Rules Service ➔ Atomic Counter ➔ HTTP 307 Redirect. Telemetry is dispatched in the background via Next.js after() without delaying the visitor.")

    # =========================================================================
    # SLIDE 7: Security and Safety
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header_badge(s7, "SECURITY & DATA INTEGRITY", width=4.5, height=0.55, left=0.8, top=0.5)

    # Subtitle
    tb_s7_sub = s7.shapes.add_textbox(Inches(5.5), Inches(0.5), Inches(7.0), Inches(0.6))
    p = tb_s7_sub.text_frame.paragraphs[0]
    p.text = "Built-in safeguards protecting visitors, privacy, and system availability"
    p.font.name = "Arial"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MUTED

    # Mascot with shield
    img_mascot_s7 = os.path.join(IMG_DIR, "Im215.png")
    if os.path.exists(img_mascot_s7):
        s7.shapes.add_picture(img_mascot_s7, Inches(0.8), Inches(1.5), width=Inches(2.5))

    # Security Safeguards List
    safeguards = [
        ("Predictable rule evaluation hierarchy (Disabled → Expired → Limit → Password)", "BUILT"),
        ("Atomic click counter prevents concurrency race conditions and overshoots", "BUILT"),
        ("Passcode protection salted and hashed with bcrypt; plain passwords never stored", "BUILT"),
        ("Stateless user sessions signed with JWT in secure httpOnly cookies (anti-XSS)", "BUILT"),
        ("Privacy-preserving visitor tracking using salted SHA-256 hashes; no raw IPs stored", "BUILT"),
        ("Password brute-force lockout: rate limits to 5 failed attempts per 15 minutes", "BUILT"),
        ("Link creation abuse protection: sliding window rate limit of 15 links per minute", "BUILT"),
        ("Automated bot exclusion: search crawlers & preview fetchers bypass click counters", "BUILT"),
        ("Admin domain blocklist: platform-wide moderation against malicious destinations", "PLANNED")
    ]

    top_pos = 1.35
    row_h = 0.47
    row_gap = 0.07

    for title, badge in safeguards:
        card = add_card(s7, 3.6, top_pos, 8.933, row_h, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=1.5)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.08)
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK

        # Badge pill
        badge_color = COLOR_GREEN_BG if badge == "BUILT" else COLOR_CYAN_LIGHT
        badge_text_color = COLOR_GREEN_TEXT if badge == "BUILT" else COLOR_CYAN
        b_pill = add_card(s7, 11.2, top_pos + 0.06, 1.2, row_h - 0.12, fill_color=badge_color, border_color=badge_text_color, border_width=1)
        p_b = b_pill.text_frame.paragraphs[0]
        p_b.text = badge
        p_b.font.name = "Arial Black"
        p_b.font.size = Pt(9)
        p_b.font.color.rgb = badge_text_color
        p_b.alignment = PP_ALIGN.CENTER

        top_pos += row_h + row_gap

    # Bottom Callout
    add_bottom_callout(s7, "Strict Tenant Isolation & Policy Control", 
        "Users only ever access their own links and analytics; administrators have global moderation oversight. A deactivated link stops traffic immediately at the first rule check while preserving complete historical logs.")

    # =========================================================================
    # SLIDE 8: Redirect Flow
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header_badge(s8, "REDIRECT FLOW", width=3.4, height=0.55, left=0.8, top=0.5)

    # Top Yellow Callout (Why 307 + no-store)
    add_card(s8, 4.4, 0.45, 8.133, 0.7, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tb_why307 = s8.shapes.add_textbox(Inches(4.4), Inches(0.48), Inches(8.133), Inches(0.65))
    p = tb_why307.text_frame.paragraphs[0]
    p.text = "Why HTTP 307 + no-store: "
    p.font.name = "Arial"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK
    run = p.add_run()
    run.text = "Permanent 301 redirects are cached by browsers, skipping our server on repeat visits. LinkPulse issues 307 + no-store so every visit evaluates current rules and records accurate telemetry."
    run.font.bold = False

    # Pipeline Steps (Horizontal)
    step_w = 2.05
    step_h = 1.3
    gap_s = 0.35
    top_step = 1.5

    steps = [
        ("Visitor Opens Short Link", "Fetches active link by indexed code from PostgreSQL", COLOR_CARD_BG, COLOR_DARK),
        ("1. Disabled?", "YES ➔ Shows branded 'Link Disabled' page with reason", RGBColor(254, 226, 226), COLOR_RED_TEXT),
        ("2. Expired?", "YES ➔ Shows 'Link Expired' page (past validity date)", RGBColor(254, 226, 226), COLOR_RED_TEXT),
        ("3. Click Limit?", "YES ➔ Shows 'Limit Reached' (checked atomically)", RGBColor(254, 226, 226), COLOR_RED_TEXT),
        ("4. Password?", "YES ➔ Shows Password Challenge screen", COLOR_YELLOW, COLOR_DARK),
    ]

    for idx, (s_title, s_desc, bg_col, txt_col) in enumerate(steps):
        s_left = 0.8 + idx * (step_w + gap_s)
        card = add_card(s8, s_left, top_step, step_w, step_h, fill_color=bg_col, border_color=COLOR_DARK, border_width=2)
        tf = card.text_frame
        tf.margin_top = Inches(0.12)
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.name = "Arial Black"
        p.font.size = Pt(10.5)
        p.font.color.rgb = txt_col
        p.alignment = PP_ALIGN.CENTER
        p_d = tf.add_paragraph()
        p_d.text = s_desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.alignment = PP_ALIGN.CENTER
        p_d.space_before = Pt(4)

    # Success Box (End of pipeline)
    succ_box = add_card(s8, 0.8 + 5 * (step_w + gap_s), top_step, step_w + 0.1, step_h, fill_color=COLOR_GREEN_BG, border_color=COLOR_GREEN_TEXT, border_width=2)
    tf_succ = succ_box.text_frame
    tf_succ.margin_top = Inches(0.2)
    p = tf_succ.paragraphs[0]
    p.text = "ALL RULES PASS"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_GREEN_TEXT
    p.alignment = PP_ALIGN.CENTER
    p_sc = tf_succ.add_paragraph()
    p_sc.text = "HTTP 307 Redirect to Destination URL"
    p_sc.font.name = "Arial"
    p_sc.font.size = Pt(9)
    p_sc.font.bold = True
    p_sc.font.color.rgb = COLOR_GREEN_TEXT
    p_sc.alignment = PP_ALIGN.CENTER

    # Rule Explanations Card
    add_card(s8, 0.8, 3.05, 11.733, 1.25, fill_color=RGBColor(241, 245, 249), border_color=COLOR_DARK, border_width=1.5)
    tb_rules_exp = s8.shapes.add_textbox(Inches(0.9), Inches(3.1), Inches(11.533), Inches(1.15))
    p = tb_rules_exp.text_frame.paragraphs[0]
    p.text = "Predictable Evaluation Hierarchy:"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK

    p_r1 = tb_rules_exp.text_frame.add_paragraph()
    p_r1.text = "• Disabled links are blocked first (protects from abusive or deleted links).  • Expired links are blocked second.  • Click limits are checked third via conditional atomic increments (preventing race condition overshoots).  • Passwords are evaluated fourth (wrong passwords keep visitors on the prompt and rate-limit failed attempts)."
    p_r1.font.name = "Arial"
    p_r1.font.size = Pt(9.5)
    p_r1.font.color.rgb = COLOR_TEXT_MUTED
    p_r1.space_before = Pt(4)

    # Bottom Background Workflow: Telemetry
    tb_bg_title = s8.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(5.0), Inches(0.5))
    p = tb_bg_title.text_frame.paragraphs[0]
    p.text = "AFTER THE REDIRECT: ASYNCHRONOUS TELEMETRY PIPELINE"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_CYAN

    telemetry_steps = [
        ("1. Bot Filtering", "Inspect User-Agent; search crawlers, preview bots & curl skip click logging"),
        ("2. Privacy Hashing", "Generate salted SHA-256 hash from IP + User-Agent; raw IP is discarded"),
        ("3. Client Classification", "Parse device family (desktop/mobile), browser, operating system & referrer"),
        ("4. Live Analytics Ingestion", "Persist ClickEvent row; immediately reflected on user & admin dashboards")
    ]
    t_w = 2.75
    t_gap = 0.244
    for idx, (t_title, t_desc) in enumerate(telemetry_steps):
        t_left = 0.8 + idx * (t_w + t_gap)
        card = add_card(s8, t_left, 5.0, t_w, 1.4, fill_color=COLOR_CARD_BG, border_color=COLOR_DARK, border_width=1.5)
        tf = card.text_frame
        tf.margin_top = Inches(0.12)
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = t_title
        p.font.name = "Arial Black"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_DARK
        p_td = tf.add_paragraph()
        p_td.text = t_desc
        p_td.font.name = "Arial"
        p_td.font.size = Pt(8.5)
        p_td.font.color.rgb = COLOR_TEXT_MUTED
        p_td.space_before = Pt(3)

    # =========================================================================
    # SLIDE 9: Data Model and Analytics
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header_badge(s9, "DATA MODEL & ANALYTICS", width=4.5, height=0.55, left=0.8, top=0.5)

    # Left: Database Entities (ER Diagram)
    add_card(s9, 0.8, 1.3, 5.6, 5.4, fill_color=RGBColor(255, 255, 255), border_color=COLOR_DARK, border_width=2)
    tb_er_title = s9.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.6), Inches(0.5))
    p = tb_er_title.text_frame.paragraphs[0]
    p.text = "RELATIONAL SCHEMA (POSTGRESQL / PRISMA)"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_CYAN
    p.alignment = PP_ALIGN.CENTER

    entities = [
        ("User", "id, email, passwordHash, role (USER | ADMIN)"),
        ("Campaign", "id, name, description, userId (deleting sets campaignId=null on links)"),
        ("Link", "id, code, destinationUrl, isActive, expiresAt, clickLimit, clicksCount, passwordHash"),
        ("ClickEvent", "id, linkId, timestamp, device, browser, os, referrer, visitorHash"),
        ("Tag & LinkTag", "Many-to-many relationship for dynamic link categorization & filtering"),
        ("PasswordAttempt", "id, linkId, ipHash, success, createdAt (salted sliding-window lockout)")
    ]

    tb_er = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.5))
    for idx, (e_name, e_fields) in enumerate(entities):
        p = tb_er.text_frame.paragraphs[0] if idx == 0 else tb_er.text_frame.add_paragraph()
        p.text = f"•  {e_name}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(7)
        run = p.add_run()
        run.text = e_fields
        run.font.bold = False
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_TEXT_MUTED

    # Right: How Clicks Become Reports
    add_card(s9, 6.7, 1.3, 5.833, 5.4, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tb_rep_title = s9.shapes.add_textbox(Inches(6.7), Inches(1.4), Inches(5.833), Inches(0.5))
    p = tb_rep_title.text_frame.paragraphs[0]
    p.text = "HOW CLICKS BECOME ACTIONABLE INSIGHTS"
    p.font.name = "Arial Black"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.alignment = PP_ALIGN.CENTER

    report_steps = [
        ("1. Bot-Filtered Event Ingestion", "Every legitimate visitor visit writes an indexed ClickEvent with device category, browser family, OS, sanitized referrer origin, and privacy-safe visitor hash."),
        ("2. Database-Level Real-Time Aggregation", "Prisma count and groupBy queries execute directly in PostgreSQL, eliminating memory bottlenecks when scaling to thousands of records."),
        ("3. Multi-Period Filtering", "Interactive time periods (7 Days, 30 Days, All Time) calculate total visits, unique visitors, active links, and daily click distributions."),
        ("4. Campaign & Audience Performance", "Aggregates click volume, link counts, and top-performing links per campaign, alongside device, browser, and referrer market shares."),
        ("5. Proactive Link Health Alerts", "Live warning cards notify owners of links expiring soon (within 72 hours), links approaching click capacity (≥ 80%), or deactivated links.")
    ]

    tb_rep = s9.shapes.add_textbox(Inches(6.9), Inches(2.0), Inches(5.4), Inches(4.5))
    for idx, (r_title, r_desc) in enumerate(report_steps):
        p = tb_rep.text_frame.paragraphs[0] if idx == 0 else tb_rep.text_frame.add_paragraph()
        p.text = f"•  {r_title}: "
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = r_desc
        run.font.bold = False
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_DARK

    # =========================================================================
    # SLIDE 10: Thank You
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)

    # Header Badge
    add_header_badge(s10, "INNOVISTA 2026 | ROUND 1: IDEATION", width=4.0, height=0.55, left=0.8, top=0.5)

    # Thank you text box
    tb_ty = s10.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(7.5), Inches(3.0))
    p = tb_ty.text_frame.paragraphs[0]
    p.text = "THANK YOU!"
    p.font.name = "Arial Black"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    p_sub = tb_ty.text_frame.add_paragraph()
    p_sub.text = "LinkPulse — URL Shortener & Telemetry Platform"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_DARK
    p_sub.space_before = Pt(8)

    p_tag = tb_ty.text_frame.add_paragraph()
    p_tag.text = "Graphic Era Hill University (GEHU), Haldwani Campus"
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(15)
    p_tag.font.color.rgb = COLOR_TEXT_MUTED
    p_tag.space_before = Pt(6)

    # Questions Welcome Box
    add_card(s10, 0.8, 4.8, 6.5, 1.2, fill_color=COLOR_YELLOW, border_color=COLOR_DARK, border_width=2)
    tb_q = s10.shapes.add_textbox(Inches(0.8), Inches(4.9), Inches(6.5), Inches(1.0))
    p = tb_q.text_frame.paragraphs[0]
    p.text = "Ready for Questions & Evaluation"
    p.font.name = "Arial Black"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_DARK
    p.alignment = PP_ALIGN.CENTER
    p_q_sub = tb_q.text_frame.add_paragraph()
    p_q_sub.text = "Live prototype and codebase available for technical demonstration"
    p_q_sub.font.name = "Arial"
    p_q_sub.font.size = Pt(10.5)
    p_q_sub.font.color.rgb = COLOR_DARK
    p_q_sub.alignment = PP_ALIGN.CENTER
    p_q_sub.space_before = Pt(4)

    # Mascot on Right
    img_mascot_s10 = os.path.join(IMG_DIR, "Im25.png")
    if os.path.exists(img_mascot_s10):
        s10.shapes.add_picture(img_mascot_s10, Inches(8.5), Inches(1.3), width=Inches(4.2))

    # Save presentation
    os.makedirs("docs", exist_ok=True)
    out_path = os.path.join("docs", "LinkPulse_Round1_Presentation.pptx")
    prs.save(out_path)
    print(f"[OK] Successfully created PowerPoint presentation at {out_path}")

if __name__ == "__main__":
    create_presentation()
